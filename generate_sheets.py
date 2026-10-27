import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import csv

data = [
    ["apaar_id", "name", "course", "seat_category", "expected_outcome_notes"],
    ["482910563721", "Aarav Sharma", "B.Tech", "General", "Accepted (Valid 12-digit, exists in registry)"],
    ["482910563721", "Aarav Sharma", "B.Tech", "General", "Rejected (Duplicate entry in file)"],
    ["12345", "Diya Patel", "B.Tech", "General", "Rejected (Invalid APAAR format: only 5 digits)"],
    ["771203948856", "Rohan Das", "B.Tech", "NRI", "Accepted (Valid APAAR, NRI quota)"],
    ["903817265540", "Meera Reddy", "B.Tech", "General", "Rejected (Not found in APAAR central registry)"],
    ["334455667788", "Vikram Malhotra", "B.Tech", "General", "Flagged (Potential dual-enrolment: Active at COEP02)"],
    ["998877665544", "Ananya Iyer", "B.Tech", "General", "Accepted (Valid APAAR, in registry)"],
    ["112233445566", "Kavita Pillai", "B.Tech", "OBC", "Accepted (Valid APAAR, in registry)"]
]

# 1. Write clean CSV file (standard format expected by portal upload)
csv_filename = "pune_college_students.csv"
with open(csv_filename, "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    for row in data:
        writer.writerow(row[:4]) # write only the 4 standard columns: apaar_id, name, course, seat_category

print(f"Generated CSV: {csv_filename}")

# 2. Write styled Excel (.xlsx) workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Student Ingestion"

# Colors
header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Indigo
header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
data_font = Font(name="Arial", size=10)
thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

for r_idx, row in enumerate(data, start=1):
    for c_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.border = thin_border
        
        if r_idx == 1:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.font = data_font
            if c_idx == 1:
                # Text alignment for APAAR ID so leading digits / precision aren't altered
                cell.number_format = '@'
                cell.alignment = Alignment(horizontal="center")
            elif c_idx == 5:
                # Color code expected notes
                if "Accepted" in val:
                    cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Light green
                elif "Rejected" in val:
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Light red
                elif "Flagged" in val:
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Light yellow

# Adjust column widths
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

excel_filename = "pune_college_students.xlsx"
wb.save(excel_filename)
print(f"Generated Excel: {excel_filename}")
